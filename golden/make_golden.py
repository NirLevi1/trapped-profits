"""יוצר תרחישי זהב: משנה תאי קלט באקסל המקורי, מחשב מחדש בליברה-אופיס, ושומר את כל תאי הנוסחאות.

הרצה:  python3 -I golden/make_golden.py
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
SRC = HERE / "logic_v1.xlsx"
OUT = HERE / "cases"

# שם תרחיש → {תא: ערך}
SCENARIOS: dict[str, dict[str, object]] = {
    "base": {},
    "transition_protected_80pct": {"חלוקות!B19": 5_760_000, "חלוקות!B20": "כן"},        # 80% → מוגן עד 2028
    "transition_35pct_year_2026": {"חלוקות!B19": 2_600_000, "חלוקות!B20": "כן"},        # 36% → מוגן ב-2026 בלבד
    "transition_not_paid": {"חלוקות!B19": 5_760_000, "חלוקות!B20": "לא"},              # לא שולם → אין הגנה
    "rule_20pct_met": {"חלוקות!B4": 700_000},                                          # 700k ≥ 20%×3.2M
    "dividend_meets_6pct": {"חלוקות!B4": 450_000},                                      # ≥ 6%×7.165M
    "dividend_meets_50pct": {"חלוקות!B4": 1_960_000},
    "losses_exempt": {"חלוקות!B5": 800_000},                                            # > 10%×7.165M
    "alt1_better_exempt_sources": {"רווחים_נצברים!B25": 6_000_000, "רווחים_נצברים!B33": 1_200_000},
    "no_few_holders": {"מבנה!B11": "לא"},
    "preferred_enterprise_no_62a": {"מבנה!B12": "כן"},
    "concentration_full_attribution": {"סעיף_62א!B37": "כן"},
    "profitability_yellow": {"סעיף_62א!B16": 1_900_000},                               # ~23%
    "profitability_green": {"סעיף_62א!B16": 1_200_000},                                # ~15%
    "exemption_750k": {"מבנה!B24": "לא", "רווחים_נצברים!B5": None, "רווחים_נצברים!B14": 600_000},
    "salary_paid_adds_to_expenses": {"מגנים!B15": "כן"},
    "balance_sheet_only": {"רווחים_נצברים!B5": None},
    "year_2025_with_avg": {"מבנה!B7": 2025, "סעיף_62א!B19": 0.22},
    "year_2025_transition_5pct": {"מבנה!B7": 2025, "חלוקות!B20": "כן"},               # 12.5% ≥ 5%
    "one_company_group": {"מבנה!B18": 1},
    "waiver_in_group": {"מבנה!B18": 3, "מבנה!B19": 1},
    "assets_shield_small": {"מגנים!B17": 3_000_000},
    "partnership_under_10pct": {"מבנה!B27": 0.05},
    "intercompany_dividend": {"סעיף_62א!B40": 400_000, "סעיף_62א!B41": 0.05},
    "zero_revenue": {"סעיף_62א!B4": 0},
}


def write_inputs(path: Path, changes: dict[str, object]) -> None:
    wb = openpyxl.load_workbook(path)
    for ref, val in changes.items():
        sheet, cell = ref.split("!")
        wb[sheet][cell].value = val
    wb.save(path)


def recalc(path: Path) -> Path:
    outdir = path.parent / "recalced"
    outdir.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return outdir / path.name


def formula_cells(src: Path) -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(src)
    out = []
    for ws in wb.worksheets:
        if ws.title == "הוראות":
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out.append((ws.title, c.coordinate))
    return out


def main() -> None:
    OUT.mkdir(exist_ok=True)
    cells = formula_cells(SRC)
    for name, changes in SCENARIOS.items():
        case_dir = OUT / name
        case_dir.mkdir(exist_ok=True)
        xlsx = case_dir / "input.xlsx"
        shutil.copy(SRC, xlsx)
        write_inputs(xlsx, changes)
        rec = recalc(xlsx)
        wb = openpyxl.load_workbook(rec, data_only=True)
        values = {f"{s}!{c}": wb[s][c].value for s, c in cells}
        (case_dir / "expected.json").write_text(json.dumps(values, ensure_ascii=False, indent=1), encoding="utf-8")
        (case_dir / "changes.json").write_text(json.dumps(changes, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{name}: {len(values)} cells", file=sys.stderr)


if __name__ == "__main__":
    main()
