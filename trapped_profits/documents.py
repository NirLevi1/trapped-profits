"""דוחות ללקוח – שני סוגים, שלושה פורמטים (החלטות 19–20).

  kind="accountant"  דוח לרו"ח: כל המספרים, שני המסלולים, הסימולטור, לבדיקה ידנית.
  kind="client"      מכתב ללקוח: מה המצב, מה מומלץ, עד מתי – בלי מונחי חוק מיותרים.

  fmt="html" | "docx" | "pdf"   עיצוב נקי כללי (אין עדיין עיצוב של המשרד – נקודה פתוחה).

מבנה ביניים אחד (Doc) → שלושה מעבדים. PDF: מ-HTML דרך דפדפן קיים (אדג'/כרום) או ליברה-אופיס; אם אין – נפתח ה-HTML להדפסה.
"""

from __future__ import annotations

import html
import os
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from .engine import ClientResult


# ---------------------------------------------------------------------------
# מבנה ביניים
# ---------------------------------------------------------------------------
@dataclass
class Block:
    kind: str                     # "h2" | "p" | "kv" | "ul" | "callout"
    text: str = ""
    rows: list[tuple[str, str, bool]] = field(default_factory=list)   # (label, value, highlight)
    items: list[str] = field(default_factory=list)
    tone: str = ""                # ל-callout: good | warn | serious | ""


@dataclass
class Doc:
    title: str
    subtitle: str
    blocks: list[Block]
    footer: str


def _ils(x: float | None) -> str:
    return "—" if x is None else f"{x:,.0f} ₪"


def _pct(x: float) -> str:
    return f"{x * 100:.1f}%"


_TONE = {"ירוק": "good", "צהוב": "warn", "אדום": "serious"}

DISCLAIMER = ("כלי ניטור וסינון ראשוני בלבד – אינו תחליף לחישוב שומתי מלא. "
              "הספרים השוטפים אינם סגורים; נתונים שנגזרו מהם הם אומדן.")


# ---------------------------------------------------------------------------
# דוח לרו"ח
# ---------------------------------------------------------------------------
def build_accountant_doc(r: ClientResult, handling: dict | None = None, history: list | None = None) -> Doc:
    d, c, s, m, sh, tr = r.dashboard, r.calc81, r.s62a, r.simulator, r.shields, r.transition
    b: list[Block] = []
    meta = [f"שנת מס {r.tax_year}", f"בסיס לתום {r.previous_year}"]
    if r.data_as_of:
        meta.append(f"נכון ל־{r.data_as_of}")
    if r.is_estimate:
        meta.append("אומדן (תרחיש ב׳)")
    meta.append(f"גרסת לוגיקה {r.logic_version}")

    b.append(Block("callout", f"סטטוס כולל: {d.status_overall}", tone=_TONE[d.color_overall]))
    b.append(Block("kv", rows=[
        ("סעיף 81 – תוספת מס 2%", d.status_81, False),
        ("סעיף 62א – רווחיות עודפת", d.status_62a, False),
    ]))
    if r.missing_fields:
        b.append(Block("callout", "חסר: " + ", ".join(r.missing_fields) + " – התוצאה חושבה עם אפס במקומם.", tone="warn"))
    if r.annualization_factor != 1.0:
        b.append(Block("p", f"נתוני השנה הוכפלו פי {r.annualization_factor:.2f} לאומדן שנתי (לינארי)."))

    b.append(Block("h2", "המלצה"))
    b.append(Block("p", m.recommendation_text or m.recommendation))

    b.append(Block("h2", "מסלול 81 – תוספת מס 2%"))
    rows = [
        (f"רווחים נצברים לתום {r.previous_year}", _ils(r.retained.retained_for_81a), False),
    ]
    if c.alternative_relevant:   # אברהם 14.10 (1): רק כשיש רווחים ממקור פטור
        rows.append(("חלופת הרווחים הפטורים (81ג(ג))", f"חלופה {c.chosen_alternative} – "
                     f"(1) מקורות פטורים {_ils(r.retained.exempt_alt1)}, (2) 7 שנים {_ils(r.retained.exempt_alt2)}", False))
    gov = {"fixed": sh.fixed_shield, "exp": sh.expenses_shield, "assets": sh.assets_shield}
    mark = lambda v: " ◀ הקובע" if abs(v - c.governing_shield) < 0.5 else ""   # noqa: E731 – אברהם 14.10 (2)
    rows += [
        ("רווחים נצברים חייבים", _ils(c.taxable_retained), False),
        ("המגן הקובע", _ils(c.governing_shield), False),
        ("— מגן קבוע" + mark(gov["fixed"]), _ils(sh.fixed_shield), bool(mark(gov["fixed"]))),
        ("— מגן הוצאות" + mark(gov["exp"]), _ils(sh.expenses_shield), bool(mark(gov["exp"]))),
        ("— מגן נכסים" + mark(gov["assets"]), _ils(sh.assets_shield), bool(mark(gov["assets"]))),
    ]
    if c.transition_reduction:
        rows.append(("הפחתה לפי הוראות המעבר", _ils(c.transition_reduction), False))
    rows += [
        ("רווחים עודפים", _ils(c.excess_profits), False),
        ("מצב", c.status_label, False),          # אברהם 14.10 (6)
        ("תוספת מס לתשלום", _ils(c.surtax), True),
    ]
    b.append(Block("kv", rows=rows))

    if m.action_required:                       # אברהם 14.10 (7): מוגן/פטור – אין סימולטור
        b.append(Block("h2", "סימולטור – לחלק דיבידנד או לשלם 2%?"))
    rows = [
        (f"{c.required_distribution_rate:.0%} מהרווחים הנצברים", _ils(m.option_6pct), False),
        ("50% מהרווחים העודפים", _ils(m.option_50pct) if m.option_50pct > 0 else "—", False),
    ]
    if m.option_20pct_start_day is not None:
        rows.append(("20% מרווחי יום התחילה", _ils(m.option_20pct_start_day), False))
    rows += [
        (f"דיבידנד מינימלי לפטור ({m.cheaper_option})", _ils(m.min_dividend), True),
        ("כבר חולק השנה", _ils(m.already_distributed), False),
        ("פער לסגירה עד 31.12", _ils(m.gap_to_close) if m.action_required else "לא נדרש", True),
        ("מס דיבידנד על הפער", _ils(m.dividend_tax_cost), False),
        ("תוספת המס אם לא מחלקים", _ils(m.surtax_if_not_distributing), False),
    ]
    if m.years_to_breakeven:
        rows.append(("שנים לאיזון", f"{m.years_to_breakeven:.1f}", False))
    if m.surtax_if_salary_paid is not None and m.surtax_if_salary_paid < c.surtax:
        rows.append(("תוספת המס אם משלמים את שכר 62א עד 31.12", _ils(m.surtax_if_salary_paid), True))
    if m.action_required:
        b.append(Block("kv", rows=rows))

    b.append(Block("h2", "מסלול 62א – רווחיות עודפת"))
    b.append(Block("kv", rows=[
        ("הכנסה מעתירת יגיעה אישית", _ils(s.personal_effort_income), False),
        ("הכנסה חייבת מפעילות זו", _ils(s.taxable_income_personal_effort), False),
        ("שיעור רווחיות", _pct(s.governing_profitability), False),
        ("חצה סף 25%?", "כן" if s.crossed_25pct else "לא", False),
        ("מתחת לסף המחזור?", "כן" if s.below_revenue_threshold else "לא", False),
        ("זכאית לפטור 750 אלף ₪?", "כן" if s.exemption_750k_applies else "לא", False),   # אברהם 14.10 (11)
        ("62א חל?", "כן" if s.section_applies else "לא", False),
        ("שכר לתיקון עד 31.12", _ils(s.attributed_income), True),
        ("שכר בעלים ששולם השנה", _ils(s.owner_salary) if s.owner_salary is not None else "לא הוזן", False),
        ("עלות מס נוספת מוערכת", _ils(s.extra_tax_cost), False),
        ("אופן החישוב", s.extra_tax_cost_method, False),
    ]))
    if s.customers_breakdown:                   # אברהם 14.10 (12)
        b.append(Block("h2", "ריכוזיות – חלוקת ההכנסות לפי לקוח"))
        b.append(Block("kv", rows=[(x["name"], f"{_ils(x['amount'])} ({x['share']:.0%})", x["share"] >= 0.55)
                                   for x in s.customers_breakdown]))
    if r.next_year_notes:                        # אברהם 14.10 (22)
        b.append(Block("h2", f"תכנון {r.tax_year + 1}"))
        b.append(Block("ul", items=list(r.next_year_notes)))

    if r.manual_review:
        b.append(Block("h2", "לבדיקה ידנית"))
        b.append(Block("ul", items=list(r.manual_review)))
    if handling and handling.get("status"):
        b.append(Block("h2", "טיפול"))
        b.append(Block("kv", rows=[("סטטוס", handling["status"], False)] +
                       ([("הערה", handling["note"], False)] if handling.get("note") else [])))
    if history:                                  # אברהם 14.10 (16, 19)
        b.append(Block("h2", "היסטוריה"))
        b.append(Block("ul", items=[f"{h['at'][:10]} · {h['kind']} · {h.get('text', '')}" for h in history[-12:][::-1]]))

    return Doc(title=f"{r.name} – רווחים כלואים", subtitle=" · ".join(meta), blocks=b, footer=DISCLAIMER)


# ---------------------------------------------------------------------------
# מכתב ללקוח
# ---------------------------------------------------------------------------
def build_client_doc(r: ClientResult, firm_name: str = "", accountant_name: str = "", letter_note: str = "") -> Doc:
    d, c, s, m, tr = r.dashboard, r.calc81, r.s62a, r.simulator, r.transition
    b: list[Block] = []
    today = date.today().strftime("%d.%m.%Y")
    b.append(Block("p", f"לכבוד {r.name}, שלום רב."))
    b.append(Block("p", f"בדקנו את מצב החברה לשנת המס {r.tax_year} מול החוק החדש על רווחים לא מחולקים "
                        f"(\"רווחים כלואים\"), שחל מתחילת 2025. להלן התמונה והמלצתנו."))

    # --- track 81 in plain words
    b.append(Block("h2", "רווחים שנצברו בחברה"))
    if not r.applies_81:
        b.append(Block("p", "החוק אינו חל על החברה. אין פעולה נדרשת."))
    elif tr.protected_in_tax_year and c.surtax == 0:
        b.append(Block("p", f"החברה חילקה דיבידנד בשנת 2025 בהיקף שמעניק הגנה עד סוף {tr.protected_until_year}. "
                            f"השנה אין צורך בחלוקה נוספת."))
    elif c.is_exempt:
        b.append(Block("p", "החברה עומדת בתנאי הפטור לשנה זו. אין תשלום נוסף ואין צורך בחלוקה נוספת."))
    elif c.excess_profits <= 0:
        b.append(Block("p", "הרווחים שנצברו בחברה נמצאים בתוך המגן שהחוק מאפשר. אין תשלום נוסף."))
    else:
        b.append(Block("p", f"בחברה נצברו רווחים של {_ils(r.retained.retained_for_81a)}. לאחר ההגנות שהחוק מאפשר, "
                            f"{_ils(c.excess_profits)} מהם חשופים לתוספת מס שנתית של 2%."))
        b.append(Block("kv", rows=[
            ("אם לא נעשה דבר – תוספת מס השנה", _ils(c.surtax), True),
            ("דיבידנד שיפטור מהתוספת (עד 31.12)", _ils(m.min_dividend), True),
            ("כבר חולק השנה", _ils(m.already_distributed), False),
            ("נותר לחלק", _ils(m.gap_to_close), True),
            ("המס על הדיבידנד הנוסף", _ils(m.dividend_tax_cost), False),
        ]))
        if m.dividend_tax_cost < c.surtax:
            b.append(Block("p", "המס על הדיבידנד נמוך מתוספת המס – חלוקה עדיפה."))
        else:
            b.append(Block("p", f"המס על הדיבידנד גבוה מתוספת המס של שנה אחת, אך תוספת המס חוזרת כל שנה"
                                + (f" ומתאזנת אחרי כ-{m.years_to_breakeven:.1f} שנים" if m.years_to_breakeven else "") + ". נשקול יחד."))

    # --- track 62a
    b.append(Block("h2", "רווחיות החברה"))
    if not r.applies_62a:
        b.append(Block("p", "הסעיף אינו חל על החברה."))
    elif r.s62a is not None and s.attributed_income > 0 and not s.section_applies:
        # מבחן הריכוזיות – ייחוס מלא בלי קשר לשיעור הרווחיות
        b.append(Block("p", f"עיקר הכנסות החברה מגיע מגורם אחד לאורך זמן, ולכן לפי החוק {_ils(s.attributed_income)} "
                            f"מהרווח מיוחסים לבעל המניות וממוסים כמשכורת – עלות נוספת מוערכת של {_ils(s.extra_tax_cost)}."))
    elif s.section_applies:
        b.append(Block("p", f"רווחיות החברה מהפעילות השוטפת היא {_pct(s.governing_profitability)}, מעל הסף של 25% שקבע החוק. "
                            f"במצב זה {_ils(s.attributed_income)} מהרווח מיוחסים לבעל המניות וממוסים כמשכורת – עלות נוספת מוערכת של {_ils(s.extra_tax_cost)}."))
        b.append(Block("p", f"ניתן למנוע זאת בתשלום שכר לבעל המניות הפעיל בסך {_ils(s.attributed_income)} עד 31.12."))
    elif s.governing_profitability > 0.20:
        b.append(Block("p", f"רווחיות החברה היא {_pct(s.governing_profitability)} – מתחת לסף 25% אך קרוב אליו. נעקוב עד סוף השנה."))
    else:
        b.append(Block("p", f"רווחיות החברה היא {_pct(s.governing_profitability)} – מתחת לסף. אין פעולה נדרשת."))

    # --- summary
    b.append(Block("h2", "לסיכום"))
    actions = []
    if m.action_required:
        actions.append(f"לחלק דיבידנד של {_ils(m.gap_to_close)} עד 31.12.{r.tax_year} (ולשלם את המס עליו בזמן).")
    if r.applies_62a and s.attributed_income > 0:
        actions.append(f"לשלם שכר לבעל המניות הפעיל בסך {_ils(s.attributed_income)} עד 31.12.{r.tax_year}.")
    if actions:
        b.append(Block("ul", items=actions))
        b.append(Block("p", "נשמח לתאם שיחה כדי להחליט יחד על הצעד המתאים."))
    else:
        b.append(Block("callout", "לא נדרשת פעולה השנה. נמשיך לעקוב.", tone="good"))
    if letter_note.strip():                      # אברהם 14.10 (18): ההערה מהשיחה
        b.append(Block("h2", "בהמשך לשיחתנו"))
        b.append(Block("p", letter_note.strip()))
    b.append(Block("p", f"בברכה,\n{accountant_name or 'רואה החשבון המטפל'}" + (f"\n{firm_name}" if firm_name else "")))

    return Doc(title=f"מצב החברה לעניין רווחים כלואים – {r.tax_year}", subtitle=f"{r.name} · {today}", blocks=b,
               footer="הערכה המבוססת על נתונים שוטפים; המספרים הסופיים ייקבעו בדוח השנתי.")


# ---------------------------------------------------------------------------
# HTML (עיצוב נקי כללי)
# ---------------------------------------------------------------------------
_CSS = """
@page{size:A4;margin:18mm 16mm}
body{font-family:'Segoe UI',Arial,'Noto Sans Hebrew',sans-serif;direction:rtl;color:#1a1a1a;font-size:11pt;line-height:1.45;max-width:720px;margin:0 auto;padding:24px 16px}
h1{font-size:20pt;font-weight:600;margin:0 0 4px}
.sub{color:#666;font-size:9.5pt;margin:0 0 18px;border-bottom:1px solid #ddd;padding-bottom:10px}
h2{font-size:12.5pt;font-weight:600;margin:20px 0 8px;color:#222}
p{margin:0 0 8px;white-space:pre-line}
table{border-collapse:collapse;width:100%;margin:4px 0 10px}
td{padding:5px 6px;border-bottom:1px solid #e6e6e6;vertical-align:top}
td.v{text-align:left;direction:ltr;white-space:nowrap;font-variant-numeric:tabular-nums}
tr.hl td{font-weight:600;background:#f6f6f6}
ul{margin:0 0 8px;padding-right:20px}
.callout{padding:9px 12px;border-radius:6px;margin:0 0 12px;font-weight:500;border-right:4px solid #888;background:#f3f3f3}
.callout.good{border-color:#12876b;background:#e6f2ee}
.callout.warn{border-color:#b0790f;background:#f7eeda}
.callout.serious{border-color:#c2542f;background:#f8e9e3}
.footer{margin-top:28px;color:#777;font-size:8.5pt;border-top:1px solid #ddd;padding-top:8px}
"""


def render_html(doc: Doc) -> str:
    e = html.escape
    out = [f"<!doctype html><html lang='he' dir='rtl'><head><meta charset='utf-8'><title>{e(doc.title)}</title><style>{_CSS}</style></head><body>",
           f"<h1>{e(doc.title)}</h1><div class='sub'>{e(doc.subtitle)}</div>"]
    for bl in doc.blocks:
        if bl.kind == "h2":
            out.append(f"<h2>{e(bl.text)}</h2>")
        elif bl.kind == "p":
            out.append(f"<p>{e(bl.text)}</p>")
        elif bl.kind == "callout":
            out.append(f"<div class='callout {bl.tone}'>{e(bl.text)}</div>")
        elif bl.kind == "ul":
            out.append("<ul>" + "".join(f"<li>{e(i)}</li>" for i in bl.items) + "</ul>")
        elif bl.kind == "kv":
            out.append("<table>" + "".join(
                f"<tr class='{'hl' if h else ''}'><td>{e(k)}</td><td class='v'>{e(v)}</td></tr>" for k, v, h in bl.rows) + "</table>")
    out.append(f"<div class='footer'>{e(doc.footer)}</div></body></html>")
    return "".join(out)


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def render_docx(doc: Doc, path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def rtl(par):
        pPr = par._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1"); pPr.append(bidi)
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return par

    def run_rtl(run, size: float | None = None, bold: bool = False, color: str | None = None):
        """טקסט עברי = complex script: גודל/הדגשה צריכים גם szCs/bCs, אחרת וורד וליברה-אופיס מתעלמים."""
        rPr = run._r.get_or_add_rPr()
        for tag, val in (("w:rtl", "1"), ("w:bCs", "1" if bold else None), ("w:b", "1" if bold else None),
                         ("w:sz", str(int(size * 2)) if size else None), ("w:szCs", str(int(size * 2)) if size else None)):
            if val is not None:
                el = OxmlElement(tag); el.set(qn("w:val"), val); rPr.append(el)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "he-IL"); lang.set(qn("w:val"), "he-IL"); rPr.append(lang)
        rFonts = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rFonts.set(qn(a), "Arial")
        rPr.insert(0, rFonts)
        return run

    d = Document()
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
    # שפת המסמך עברית (אחרת וורד מנחש "ערבית") ומצב תאימות של וורד עדכני (בלי "[Compatibility Mode]" בכותרת)
    st_rPr = st.element.get_or_add_rPr()
    lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "he-IL"); lang.set(qn("w:bidi"), "he-IL"); st_rPr.append(lang)
    settings = d.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat"); settings.append(compat)
    for cs in list(compat):
        if cs.tag == qn("w:compatSetting") and cs.get(qn("w:name")) == "compatibilityMode":
            compat.remove(cs)
    cs = OxmlElement("w:compatSetting")
    cs.set(qn("w:name"), "compatibilityMode"); cs.set(qn("w:uri"), "http://schemas.microsoft.com/office/word"); cs.set(qn("w:val"), "15")
    compat.append(cs)
    for s in d.sections:
        s.left_margin = s.right_margin = Pt(50)

    p = rtl(d.add_paragraph()); run_rtl(p.add_run(doc.title), size=18, bold=True)
    p = rtl(d.add_paragraph()); run_rtl(p.add_run(doc.subtitle), size=9, color="666666")

    for bl in doc.blocks:
        if bl.kind == "h2":
            p = rtl(d.add_paragraph()); p.paragraph_format.space_before = Pt(12)
            run_rtl(p.add_run(bl.text), size=13, bold=True)
        elif bl.kind in ("p", "callout"):
            for line in bl.text.split("\n"):
                p = rtl(d.add_paragraph()); run_rtl(p.add_run(line), bold=(bl.kind == "callout"))
                if bl.kind == "callout":
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), {"good": "E6F2EE", "warn": "F7EEDA", "serious": "F8E9E3"}.get(bl.tone, "F3F3F3"))
                    p._p.get_or_add_pPr().append(shd)
        elif bl.kind == "ul":
            for i in bl.items:
                p = rtl(d.add_paragraph(style="List Bullet")); run_rtl(p.add_run(i))
        elif bl.kind == "kv":
            t = d.add_table(rows=0, cols=2); t.style = "Table Grid"
            tblPr = t._tbl.tblPr
            bidi = OxmlElement("w:bidiVisual"); tblPr.append(bidi)
            for k, v, h in bl.rows:
                cells = t.add_row().cells
                for cell, txt in ((cells[0], k), (cells[1], v)):
                    cell.text = ""
                    par = rtl(cell.paragraphs[0]); run_rtl(par.add_run(txt), bold=h)
            d.add_paragraph()

    p = rtl(d.add_paragraph()); run_rtl(p.add_run(doc.footer), size=8, color="777777")
    d.save(str(path))
    return path


# ---------------------------------------------------------------------------
# PDF – מ-HTML דרך דפדפן/ליברה-אופיס שקיימים במחשב
# ---------------------------------------------------------------------------
def _find_browser() -> str | None:
    cands = []
    if sys.platform.startswith("win"):
        pf = [os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
              os.environ.get("LocalAppData", "")]
        for base in pf:
            cands += [os.path.join(base, r"Microsoft\Edge\Application\msedge.exe"),
                      os.path.join(base, r"Google\Chrome\Application\chrome.exe")]
    if sys.platform == "darwin":
        cands += ["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                  "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
                  "/Applications/Chromium.app/Contents/MacOS/Chromium",
                  "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser",
                  os.path.expanduser("~/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")]
    for name in ("msedge", "chrome", "google-chrome", "chromium", "chromium-browser"):
        if shutil.which(name):
            cands.append(shutil.which(name))
    return next((c for c in cands if c and os.path.exists(c)), None)


def html_to_pdf(html_path: Path, pdf_path: Path, doc: "Doc | None" = None) -> Path | None:
    browser = _find_browser()
    if browser:
        try:
            with tempfile.TemporaryDirectory() as profile:   # פרופיל נפרד – עובד גם כשהדפדפן כבר פתוח
                subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-sandbox", f"--user-data-dir={profile}",
                                f"--print-to-pdf={pdf_path}", "--no-pdf-header-footer", html_path.resolve().as_uri()],
                               check=True, capture_output=True, timeout=60)
            if pdf_path.exists():
                return pdf_path
        except Exception:  # noqa: BLE001
            pass
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        # ליברה-אופיס ממיר וורד הרבה יותר טוב מ-HTML (צבעים, טבלאות, כיוון)
        src = html_path
        if doc is not None:
            src = render_docx(doc, pdf_path.with_suffix(".tmp.docx"))
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(src)],
                           check=True, capture_output=True, timeout=120)
            produced = pdf_path.parent / (src.stem + ".pdf")
            if produced.exists():
                if produced != pdf_path:
                    produced.replace(pdf_path)
                return pdf_path
        except Exception:  # noqa: BLE001
            pass
        finally:
            if src != html_path and src.exists():
                src.unlink()
    return None


# ---------------------------------------------------------------------------
# נקודת כניסה
# ---------------------------------------------------------------------------
def export(r: ClientResult, kind: str, fmt: str, out_dir: Path, handling: dict | None = None,
           letter_note: str = "", history: list | None = None,
           firm_name: str = "", accountant_name: str = "") -> Path | None:
    doc = build_client_doc(r, firm_name, accountant_name, letter_note) if kind == "client" else build_accountant_doc(r, handling, history)
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = "".join(ch for ch in r.name if ch not in '\\/:*?"<>|').strip() or "client"
    stem = f"{safe} – {'מכתב ללקוח' if kind == 'client' else 'דוח לרואה חשבון'} {r.tax_year}"
    if fmt == "docx":
        return render_docx(doc, out_dir / f"{stem}.docx")
    html_path = out_dir / f"{stem}.html"
    html_path.write_text(render_html(doc), encoding="utf-8")
    if fmt == "html":
        return html_path
    if fmt == "pdf":
        return html_to_pdf(html_path, out_dir / f"{stem}.pdf", doc) or html_path   # אין ממיר → נפתח ה-HTML להדפסה
    raise ValueError(fmt)
